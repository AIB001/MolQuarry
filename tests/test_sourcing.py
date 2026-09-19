import json

import pytest

from molquarry.errors import MolQuarryError
from molquarry.models import QueryResult
from molquarry.workflows.journal import QueryJournal
from molquarry.workflows.sourcing import (
    SourcingConfig,
    local_catalog_matches,
    lookup_batches,
    read_sdf,
    screen_sdf,
)

Chem = pytest.importorskip("rdkit.Chem")
pytest.importorskip("openpyxl")


def sdf(path, smiles=("CCO", "CCO", "C[C@H](O)C(=O)O")):
    with Chem.SDWriter(str(path)) as writer:
        for index, value in enumerate(smiles):
            molecule = Chem.MolFromSmiles(value)
            molecule.SetProp("_Name", f"record-{index}")
            if "@" in value:
                molecule.SetProp("MOLCASCADE_PARENT_SMILES", value.replace("@", ""))
            writer.write(molecule)


class FakeQuarry:
    def __init__(self, fail=False):
        self.calls = []
        self.fail = fail

    def query(self, source, operation, **params):
        self.calls.append((source, operation, params))
        if self.fail:
            raise MolQuarryError("rate_limited", "fixture")
        records = []
        if source == "pubchem" and operation == "batch_properties":
            ethanol = Chem.MolToInchiKey(Chem.MolFromSmiles("CCO"))
            if ethanol in params["inchikeys"]:
                records = [{"CID": 702, "InChIKey": ethanol}]
        if operation == "source_categories":
            records = [
                {
                    "Category": "Chemical Vendors",
                    "Sources": [{"SourceName": "Fixture supplier", "RegistryID": "A1"}],
                }
            ]
        return QueryResult(
            source=source,
            operation=operation,
            parameters=params,
            records=records,
            returned=len(records),
            provenance=[],
        )


def test_records_duplicates_and_stereo_are_preserved(tmp_path):
    path = tmp_path / "input.sdf"
    sdf(path)
    rows, keys = read_sdf(path, SourcingConfig())
    assert len(rows) == 3 and len(keys) == 3
    assert rows[0]["inchikey"] == rows[1]["inchikey"]
    assert rows[2]["parent_identity_status"] == "stereo_or_other_layer_differs"
    with pytest.raises(ValueError, match="size budget"):
        read_sdf(path, SourcingConfig(max_input_bytes=1))
    with pytest.raises(ValueError, match="no truncated"):
        read_sdf(path, SourcingConfig(max_records=1))


def test_partial_batch_matches_by_key_and_failure_is_not_absence(tmp_path):
    keys = [Chem.MolToInchiKey(Chem.MolFromSmiles(s)) for s in ["CCN", "CCO"]]
    journal = QueryJournal(FakeQuarry(), tmp_path / "success")
    found, statuses = lookup_batches(journal, "pubchem", "batch_properties", keys, 100, print)
    assert list(found) == [keys[1]]
    assert statuses[keys[0]]["status"] == "no_exact_match_in_source"
    failed = QueryJournal(FakeQuarry(fail=True), tmp_path / "failure")
    found, statuses = lookup_batches(failed, "pubchem", "batch_properties", keys, 100, print)
    assert not found and {s["status"] for s in statuses.values()} == {"rate_limited"}


def test_screen_workbook_resume_and_unknown_synthesis(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "input.sdf"
    sdf(path)
    quarry = FakeQuarry()
    config = SourcingConfig(unichem=False, max_mcule_queries=0)
    output = tmp_path / "screen"
    result = screen_sdf(quarry, path, output, config=config)
    assert result["input_records"] == 3 and result["unique_sdf_identities"] == 2
    assert result["records_with_vendor_listings"] == 2
    assert result["confirmed_in_stock"] == result["validated_synthesis_routes"] == 0
    book = load_workbook(output / result["workbook"], read_only=True)
    assert len(book.sheetnames) == 9 and book["Compounds"].max_row == 4
    book.close()
    before = len(quarry.calls)
    screen_sdf(quarry, path, output, config=config, resume=True)
    assert len(quarry.calls) == before
    with pytest.raises(ValueError, match="fingerprint differs"):
        screen_sdf(quarry, path, output, config=SourcingConfig(), resume=True)


def test_disabled_sources_are_unsearched(tmp_path):
    path = tmp_path / "input.sdf"
    sdf(path, ["CCO"])
    output = tmp_path / "screen"
    config = SourcingConfig(pubchem=False, chembl=False, unichem=False, max_mcule_queries=0)
    screen_sdf(FakeQuarry(), path, output, config=config)
    result = json.loads((output / "results.json").read_text())
    assert result["compounds"][0]["existence_status"] == "not_searched"


def test_journal_retries_errors_only_when_requested(tmp_path):
    quarry = FakeQuarry(fail=True)
    journal = QueryJournal(quarry, tmp_path)
    assert journal.query("pubchem", "batch_properties", inchikeys=[])["status"] == "rate_limited"
    quarry.fail = False
    assert journal.query("pubchem", "batch_properties", inchikeys=[])["status"] == "rate_limited"
    journal.retry_errors = True
    assert journal.query("pubchem", "batch_properties", inchikeys=[])["status"] == "ok"
    assert len(quarry.calls) == 2


def test_repeated_batch_continuation_is_incomplete_not_absent(tmp_path):
    class Repeating(FakeQuarry):
        def query(self, source, operation, **params):
            result = super().query(source, operation, **params)
            result.next_parameters = params
            return result

    journal = QueryJournal(Repeating(), tmp_path)
    _, statuses = lookup_batches(journal, "chembl", "batch_details", ["key"], 50, print)
    assert statuses["key"]["status"] == "partial_pagination"


def test_resume_refuses_a_changed_archived_input(tmp_path):
    path = tmp_path / "input.sdf"
    sdf(path, ["CCO"])
    output = tmp_path / "run"
    config = SourcingConfig(pubchem=False, chembl=False, unichem=False, max_mcule_queries=0)
    screen_sdf(FakeQuarry(), path, output, config=config)
    (output / "input/shortlist.sdf").write_text("changed")
    with pytest.raises(ValueError, match="Saved input copy"):
        screen_sdf(FakeQuarry(), path, output, config=config, resume=True)


def test_local_catalog_wrong_column_count_does_not_become_no_match(tmp_path):
    path = tmp_path / "catalog.csv"
    path.write_text("name,smiles\none,CCO,unexpected\n")
    with pytest.raises(ValueError, match="malformed row"):
        local_catalog_matches([str(path)], {})


def test_identity_role_can_be_both_sdf_and_parent_variant(tmp_path):
    path = tmp_path / "input.sdf"
    sdf(path, ["C[C@H](O)C(=O)O", "CC(O)C(=O)O"])
    rows, identities = read_sdf(path, SourcingConfig())
    assert identities[rows[1]["inchikey"]]["role"] == "sdf_and_parent_variant"
